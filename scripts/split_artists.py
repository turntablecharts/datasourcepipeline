"""Expand a CSV or XLSX song list to one CSV row per song, artist, and year.

Usage:
    python scripts/split_artists.py songs.xlsx artists.csv
    python scripts/split_artists.py songs.csv artists.csv
"""

import argparse
import csv
from datetime import date, datetime
import html
from pathlib import Path
import re


def clean(value):
    text = html.unescape(str(value)) if value is not None else ""
    return " ".join(re.sub(r"<br\s*/?>", " ", text, flags=re.I).split())


def expand_rows(rows, title_column, artist_column, date_column):
    output = []
    seen = set()
    for number, raw in enumerate(rows, start=2):
        row = {clean(key): value for key, value in raw.items()}
        title = clean(row.get(title_column))
        artists = clean(row.get(artist_column))
        entry_date = row.get(date_column)
        if not any((title, artists, clean(entry_date))):
            continue
        if not title or not artists:
            raise ValueError(f"Row {number}: missing song title or artist name")
        if isinstance(entry_date, (date, datetime)):
            year = entry_date.year
        else:
            years = re.findall(r"\b(?:19|20)\d{2}\b", clean(entry_date))
            if len(set(years)) != 1:
                raise ValueError(f"Row {number}: expected one four-digit entry year")
            year = int(years[0])
        names = [clean(name) for name in re.split(r"[,&]", artists) if clean(name)]
        if not names:
            raise ValueError(f"Row {number}: no artist names found")
        for name in names:
            key = (title.casefold(), name.casefold(), year)
            if key not in seen:
                seen.add(key)
                output.append({"Title": title, "Artiste Name": name, "Year": year})
    return output


def read_rows(path):
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as source:
            reader = csv.DictReader(source)
            return list(reader.fieldnames or []), list(reader)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            values = workbook.worksheets[0].iter_rows(values_only=True)
            headers = next(values, ())
            return list(headers), [dict(zip(headers, row)) for row in values]
        finally:
            workbook.close()
    raise ValueError("Input must be a .csv or .xlsx file")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path, help="New output CSV file")
    parser.add_argument("--title-column", default="Title")
    parser.add_argument("--artist-column", default="Artiste Name")
    parser.add_argument("--date-column", default="Date of First Entry at No. 1")
    args = parser.parse_args()
    try:
        if args.output.suffix.lower() != ".csv":
            raise ValueError("Output must be a .csv file")
        headers, rows = read_rows(args.input)
        required = [args.title_column, args.artist_column, args.date_column]
        missing = set(required) - {clean(header) for header in headers}
        if missing:
            raise ValueError(f"Missing columns: {', '.join(sorted(missing))}")
        output = expand_rows(rows, *required)
        with args.output.open("x", encoding="utf-8-sig", newline="") as destination:
            writer = csv.DictWriter(destination, fieldnames=["Title", "Artiste Name", "Year"])
            writer.writeheader()
            writer.writerows(output)
    except (ValueError, OSError) as exc:
        parser.exit(1, f"Error: {exc}\n")
    print(f"Wrote {len(output)} song/artist/year rows to {args.output}")


if __name__ == "__main__":
    main()
