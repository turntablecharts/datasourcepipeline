from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import and_, func, select
from sqlalchemy.orm import Session

from src.models import AudiomackStream, BoomplayStream


def aggregate_rows(rows, limit: int) -> list[dict]:
    grouped: dict[tuple[str, str], dict] = {}
    for row in rows:
        key = (row.artist_normalized, row.song_title_normalized)
        current = grouped.setdefault(key, {"artist": row.artist, "song": row.song_title, "streams": 0, "date": row.play_date})
        current["streams"] += row.streams
        if row.play_date >= current["date"]:
            current.update(artist=row.artist, song=row.song_title, date=row.play_date)
    result = sorted(grouped.values(), key=lambda item: (-item["streams"], item["artist"].casefold(), item["song"].casefold()))
    if len(result) > limit:
        cutoff = result[limit - 1]["streams"]
        result = [item for item in result if item["streams"] >= cutoff]
    return [{"Artist": item["artist"], "Song": item["song"], "Streams": item["streams"]} for item in result]


def _aggregated_query(model, start: date, end: date, limit: int, countries: list[str] | None = None):
    has_country = model is AudiomackStream
    group_columns = [model.artist_normalized, model.song_title_normalized]
    filters = [model.play_date.between(start, end)]
    if has_country:
        group_columns.insert(0, model.country_code)
        filters.append(model.country_code.in_(countries or []))

    aggregate_columns = [
        *( [model.country_code.label("country_code")] if has_country else [] ),
        model.artist_normalized.label("artist_normalized"),
        model.song_title_normalized.label("song_title_normalized"),
        func.sum(model.streams).label("total_streams"),
    ]
    totals = (
        select(*aggregate_columns)
        .where(*filters)
        .group_by(*group_columns)
        .subquery("stream_totals")
    )

    display_partition = [model.artist_normalized, model.song_title_normalized]
    if has_country:
        display_partition.insert(0, model.country_code)
    display_columns = [
        *( [model.country_code.label("country_code")] if has_country else [] ),
        model.artist_normalized.label("artist_normalized"),
        model.song_title_normalized.label("song_title_normalized"),
        model.artist.label("artist"),
        model.song_title.label("song_title"),
        func.row_number().over(
            partition_by=display_partition,
            order_by=(model.play_date.desc(), model.created_at.desc(), model.id.desc()),
        ).label("display_rank"),
    ]
    latest_spelling = select(*display_columns).where(*filters).subquery("latest_spelling")
    display = select(latest_spelling).where(latest_spelling.c.display_rank == 1).subquery("display")

    join_conditions = [
        totals.c.artist_normalized == display.c.artist_normalized,
        totals.c.song_title_normalized == display.c.song_title_normalized,
    ]
    if has_country:
        join_conditions.append(totals.c.country_code == display.c.country_code)
    combined_columns = [
        *( [totals.c.country_code] if has_country else [] ),
        display.c.artist,
        display.c.song_title,
        totals.c.total_streams,
    ]
    combined = select(*combined_columns).join(display, and_(*join_conditions)).subquery("combined")

    rank_args = {"order_by": combined.c.total_streams.desc()}
    if has_country:
        rank_args["partition_by"] = combined.c.country_code
    ranked = select(
        *combined.c,
        func.rank().over(**rank_args).label("stream_rank"),
    ).subquery("ranked")

    order_columns = [ranked.c.total_streams.desc(), ranked.c.artist.asc(), ranked.c.song_title.asc()]
    if has_country:
        order_columns.insert(0, ranked.c.country_code.asc())
    return select(ranked).where(ranked.c.stream_rank <= limit).order_by(*order_columns)


def query_boomplay(db: Session, start: date, end: date, limit: int = 500) -> list[dict]:
    rows = db.execute(_aggregated_query(BoomplayStream, start, end, limit)).mappings().all()
    return [{"Artist": row["artist"], "Song": row["song_title"], "Streams": row["total_streams"]} for row in rows]


def query_audiomack(
    db: Session, start: date, end: date, countries: list[str], limit: int = 800
) -> dict[str, list[dict]]:
    normalized = sorted(set(country.upper() for country in countries))
    rows = db.execute(
        _aggregated_query(AudiomackStream, start, end, limit, normalized)
    ).mappings().all()
    by_country = {country: [] for country in normalized}
    for row in rows:
        by_country[row["country_code"]].append({
            "Artist": row["artist"],
            "Song": row["song_title"],
            "Streams": row["total_streams"],
        })
    return by_country


def build_workbook(sheets: dict[str, list[dict]]) -> BytesIO:
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="1D4ED8")
    for sheet_name, rows in sheets.items():
        sheet = workbook.create_sheet(title=sheet_name[:31])
        sheet.append(["Artist", "Song", "Streams"])
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
        for row in rows:
            sheet.append([row["Artist"], row["Song"], row["Streams"]])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:C{max(sheet.max_row, 1)}"
        sheet.column_dimensions["A"].width = 32
        sheet.column_dimensions["B"].width = 42
        sheet.column_dimensions["C"].width = 16
        for cell in sheet["C"][1:]:
            cell.number_format = "#,##0"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
